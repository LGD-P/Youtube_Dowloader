import docx
import openpyxl
import odf.opendocument
from odf import teletype
from odf.opendocument import load
import pandas as pd



def read_txt(file_path):
    print("For TXT")
    with open(file_path, "r", encoding='utf-8') as f:
        result = [line.strip() for line in f if line.strip()]
        print(result)
        return result


def read_ods(file_path):
    print("For Libre Office Calc")
    result = []
    datas = pd.read_excel(file_path, engine="odf", header=None)
    for _, row in datas.iterrows():
        data = []
        for value in row:
            data.append(value)
        for e in data:
            if not pd.isnull(e):
                result.append(data[0])
    return result


def read_odt(file_path):
    print("For Libre Office TXT")
    textdoc = load(file_path)
    first_step = [teletype.extractText(para) for para in textdoc.getElementsByType(odf.text.P)]
    result = [e for e in first_step if e != '']
    return result


def read_docx(file_path):
    print("For WORD TXT")
    doc = docx.Document(file_path)
    result = [para.text for para in doc.paragraphs if para.text.strip()]
    print(result)
    return result


def read_xlsx(file_path):
    print("For EXCEL")
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    result = [cell.value for row in sheet.iter_rows()
              for cell in row if cell.value]
    print(result)
    return result


def read_any_file_and_return_list(file_path):
    """Mapp reading function to each file type
    return a list after reading
    """
    file_readers = {
        ".txt": read_txt,
        ".odt": read_odt,
        ".ods": read_ods,
        ".docx": read_docx,
        ".xlsx": read_xlsx
    }

    for extension, reader in file_readers.items():
        if extension in file_path:
            return reader(file_path)

    return print("Error: Unsupported file format")

